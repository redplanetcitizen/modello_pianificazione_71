"""Passo C2: margini dell'investimento.

Confronta, per anno e tipo di bene, l'investimento privato dei Fixed Assets dettagliati
(per industria e tipo di bene, prezzi correnti) con le colonne d'investimento della tavola Use
(F02E attrezzature, F02S strutture, F02N proprietà intellettuale, F02R residenziale) e con la NIPA 5.3.5.

Serve a stabilire, PRIMA di ogni bilanciamento (RAS/GRAS), se i due margini della matrice Φ
hanno lo stesso perimetro, la stessa valutazione e lo stesso totale (specifica §5.2).
Costruisce anche l'investimento per industria I/O e tipo, tramite la concordanza FA 74 → I/O (H18).
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from .sistema import ANNI

TIPI_FA = {"E": "EQ00", "S": "ST00", "N": "IP00"}
COLONNE_USE = {"E": "F02E", "S": "F02S", "N": "F02N", "R": "F02R"}
# NIPA 5.3.5: righe di riferimento
RIGHE_NIPA = {"E": 9, "S": 3, "N": 16, "R": 20, "S_nuove": 29}


def leggi_fa_dettaglio(percorso: Path | str) -> pd.DataFrame:
    """Foglio 'Datasets' dei file detailnonres_*: serie I3N<industria>1<bene>.A per anno."""
    wb = openpyxl.load_workbook(percorso, read_only=True, data_only=True)
    righe = list(wb["Datasets"].iter_rows(values_only=True))
    wb.close()
    anni = [str(a) for a in righe[0][1:]]
    df = pd.DataFrame([r for r in righe[1:] if r[0]], columns=["serie"] + anni)
    if not df["serie"].str.fullmatch(r"I3N.{4}1.{4}\.A").all():
        raise ValueError("Codici di serie inattesi nel foglio Datasets")
    df["industria_fa"] = df["serie"].str[3:7]
    df["bene"] = df["serie"].str[8:12]
    lungo = df.melt(id_vars=["serie", "industria_fa", "bene"], value_vars=[str(a) for a in ANNI],
                    var_name="anno", value_name="valore")
    lungo["anno"] = lungo["anno"].astype(int)
    lungo["valore"] = pd.to_numeric(lungo["valore"], errors="raise").astype(float)
    return lungo


def leggi_residenziale(percorso: Path | str, foglio: str = "Investment ") -> pd.Series:
    """Totale dell'investimento residenziale privato (riga 'RESIDENTIAL TOTALS'), per anno."""
    wb = openpyxl.load_workbook(percorso, read_only=True, data_only=True)
    righe = list(wb[foglio].iter_rows(values_only=True))
    wb.close()
    intest = next(r for r in righe if r and r[0] == "Asset Codes")
    tot = next(r for r in righe if len(r) > 1 and isinstance(r[1], str) and r[1].strip() == "RESIDENTIAL TOTALS")
    colonne = {str(v): i for i, v in enumerate(intest) if v is not None}
    return pd.Series({a: float(tot[colonne[str(a)]]) for a in ANNI}, name="FA_residenziale")


def leggi_concordanza(percorso: Path | str) -> pd.DataFrame:
    c = pd.read_csv(percorso, dtype=str)
    if c["industria_fa"].duplicated().any():
        raise ValueError("Concordanza FA → I/O: industria FA duplicata")
    return c


def confronto_totali(fa: pd.DataFrame, res_fa: pd.Series, use_f: dict[int, pd.DataFrame],
                     nipa: pd.DataFrame) -> pd.DataFrame:
    """Totale per tipo e anno: Fixed Assets, Use (con e senza la riga Used), NIPA 5.3.5."""
    righe = []
    n = nipa[nipa["TimePeriod"].astype(str).isin([str(a) for a in ANNI])].copy()
    n["anno"] = n["TimePeriod"].astype(int)
    for anno in ANNI:
        F = use_f[anno]
        for tipo in ("E", "S", "N", "R"):
            col = COLONNE_USE[tipo]
            fa_tot = res_fa[anno] if tipo == "R" else fa[(fa["anno"] == anno) & (fa["bene"] == TIPI_FA[tipo])]["valore"].sum()
            use_tot = F[col].sum()
            use_senza_used = F[col].drop(index="Used").sum()
            nipa_tot = n[(n["anno"] == anno) & (n["LineNumber"] == RIGHE_NIPA[tipo])]["value_num"].sum()
            riga = {"anno": anno, "tipo": tipo, "fixed_assets": fa_tot, "use_totale": use_tot,
                    "use_riga_used": F.loc["Used", col], "use_senza_used": use_senza_used, "nipa_535": nipa_tot,
                    "scarto_fa_use": fa_tot - use_tot, "scarto_relativo": (fa_tot - use_tot) / use_tot,
                    "use_meno_nipa": use_tot - nipa_tot}
            if tipo == "S":
                riga["nipa_535_strutture_nuove"] = n[(n["anno"] == anno) & (n["LineNumber"] == RIGHE_NIPA["S_nuove"])]["value_num"].sum()
            righe.append(riga)
    return pd.DataFrame(righe)


def investimento_per_industria_io(fa: pd.DataFrame, conc: pd.DataFrame) -> pd.DataFrame:
    """Investimento per industria I/O, tipo e anno (prezzi correnti), dalla concordanza FA → I/O."""
    mancanti = sorted(set(fa["industria_fa"]) - set(conc["industria_fa"]))
    if mancanti:
        raise ValueError(f"Industrie FA senza corrispondenza I/O: {mancanti}")
    tot = fa[fa["bene"].isin(TIPI_FA.values())].merge(conc[["industria_fa", "industria_io"]], on="industria_fa")
    tot["tipo"] = tot["bene"].map({v: k for k, v in TIPI_FA.items()})
    return tot.groupby(["anno", "industria_io", "tipo"], as_index=False)["valore"].sum()
