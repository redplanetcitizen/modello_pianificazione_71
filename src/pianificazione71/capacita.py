"""Passo C6: capacità produttiva legata al capitale (specifica §4.3, §6.3).

1. Costo d'uso per industria e tipo (metodo BLS/Jorgenson, senza guadagni in conto capitale):
       c_{j,a} = r_g + δ_{j,a},
   con r_g il tasso di rendimento del gruppo KLEMS g che contiene j, calibrato nel 2012 in modo che
       Σ_{j∈g,a} (r_g + δ_{j,a}) K_{j,a} = remunerazione del capitale KLEMS del gruppo.
   La remunerazione KLEMS comprende anche terreni e scorte, che non sono in K: r è sovrastimato (H8b).
2. Pesi dell'aggregato lineare (H8): w_{j,a} = c_{j,a} · ΣK_j / Σ_a c_{j,a} K_{j,a}, così che nel 2012
   K^cap_j = Σ_a w_{j,a} K_{j,a} = Σ_a K_{j,a} (unità: dollari di stock 2012 a composizione di servizi).
3. κ_j = K^cap_{j,2012} · u_{j,2012} / x_{j,2012}; u dalla Fed G.17 dove esiste, altrimenti 1 (H9).
4. Controllo esterno parziale (§6.3): utilizzo implicito u*_{j,t} = x_{j,t} κ_j / K^cap_{j,t} contro u G.17.
"""
from __future__ import annotations

import re

import numpy as np
import openpyxl
import pandas as pd

CATEGORIE_CAPITALE = ("Capital_Art Compensation", "Capital_IT Compensation", "Capital_Other Compensation",
                      "Capital_R&D Compensation", "Capital_Software Compensation")
KLEMS_AGGREGATI = {"44RT": ["441", "445", "452", "4A0"], "531": ["HS", "ORE"], "622HO": ["622", "623"],
                   "GF": ["GFGD", "GFGN", "GFE"], "GSL": ["GSLG", "GSLE"]}

# Serie G.17 (tavole 7 e 8) → industrie I/O. Gruppi "fisici": capacità stimata soprattutto da dati fisici
# (Fed, MethCap: estrazione, utility, carta, raffinazione, metalli di base, autoveicoli; chimica in parte).
G17_IO = {
    "G321": ["321"], "G327": ["327"], "G331": ["331"], "G332": ["332"], "G333": ["333"], "G334": ["334"],
    "G335": ["335"], "G3361T3": ["3361MV"], "G3364T9": ["3364OT"], "G337": ["337"], "G339": ["339"],
    "G311A2": ["311FT"], "G313A4": ["313TT"], "G315A6": ["315AL"], "G322": ["322"], "G323": ["323"],
    "G324": ["324"], "G325": ["325"], "G326": ["326"], "G21": ["211", "212", "213"], "G2211A2": ["22"],
}
METODO_G17 = {"G21": "fisico", "G2211A2": "fisico", "G322": "fisico", "G324": "fisico", "G331": "fisico",
              "G3361T3": "fisico", "G325": "misto"}  # tutti gli altri: indagine Census + capitale


# Etichette dei fogli dati che differiscono da quelle del foglio "NAICS codes" (stessa industria, stessa posizione).
KLEMS_ALIAS = {
    "Publishing industries, except internet (includes software)": "Publishing industries (includes software)",
    "Data processing, internet publishing, and other information services": "Information and data processing services",
    "Federal": "Federal government",
    "State and local": "State and local government",
}


def _righe_dati_klems(righe) -> list:
    """Righe di dati di un foglio KLEMS: etichetta nella colonna A e almeno un valore numerico negli anni.

    Esclude intestazioni, righe vuote e note a piè di foglio (in alcuni fogli, per esempio
    `Integrated Labor Productivity` e `Integrated TFP Index`, "Note:" e la nota con l'asterisco sono
    testo nella colonna A: contarle come righe dava 65 righe contro 63 codici).
    """
    return [r for r in righe[2:] if r[0] not in (None, "")
            and any(isinstance(v, (int, float)) and not isinstance(v, bool) for v in r[1:])]


def leggi_klems(percorso, fogli, anni) -> pd.DataFrame:
    """Valori KLEMS per codice di industria (63) e anno, somma dei fogli indicati.

    Le righe si associano ai codici per posizione, dopo aver verificato che ogni etichetta coincida
    con la descrizione del foglio "NAICS codes" nella stessa posizione (a meno degli alias dichiarati).
    """
    wb = openpyxl.load_workbook(percorso, read_only=True, data_only=True)
    tab = [r for r in list(wb["NAICS codes"].iter_rows(values_only=True))[1:]
           if r[1] not in (None, "") and str(r[1]).strip() != "Production Account Codes"]
    codici = [str(r[1]).strip() for r in tab]
    descrizioni = [str(r[0]).strip() for r in tab]
    tot = None
    for f in fogli:
        righe = list(wb[f].iter_rows(values_only=True))
        intest = [str(v) for v in righe[1]]
        dati = _righe_dati_klems(righe)
        if len(dati) != len(codici):
            raise ValueError(f"KLEMS {f}: {len(dati)} righe di dati contro {len(codici)} codici")
        for k, r in enumerate(dati):
            etichetta = str(r[0]).strip()
            if KLEMS_ALIAS.get(etichetta, etichetta) != descrizioni[k]:
                raise ValueError(f"KLEMS {f}, riga {k + 1}: '{etichetta}' non corrisponde a '{descrizioni[k]}' ({codici[k]})")
        df = pd.DataFrame([[r[intest.index(str(a))] for a in anni] for r in dati], index=codici, columns=list(anni))
        df = df.apply(pd.to_numeric, errors="raise").astype(float)
        tot = df if tot is None else tot + df
    wb.close()
    return tot


def gruppo_klems(industria_io: str) -> str:
    for g, membri in KLEMS_AGGREGATI.items():
        if industria_io in membri:
            return g
    return industria_io


def pesi_costo_uso(pannello: pd.DataFrame, remunerazione: pd.Series, anno: int = 2012) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Restituisce (pesi per industria e tipo, tassi di rendimento per gruppo KLEMS)."""
    p = pannello[(pannello["anno"] == anno) & pannello["delta"].notna()].copy()
    p = p[~p["industria_io"].str.startswith("G")]  # settore pubblico esogeno (H23)
    p["gruppo"] = p["industria_io"].map(gruppo_klems)
    p["DK"] = p["delta"] * p["K_inizio_2012"]
    g = p.groupby("gruppo").agg(K=("K_inizio_2012", "sum"), DK=("DK", "sum"))
    g["remunerazione"] = remunerazione.reindex(g.index)
    g["r"] = (g["remunerazione"] - g["DK"]) / g["K"]
    g["r_negativo_troncato"] = g["r"] < 0
    g["r_usato"] = g["r"].clip(lower=0)
    p = p.join(g["r_usato"], on="gruppo")
    p["costo_uso"] = p["r_usato"] + p["delta"]
    somma_cK = (p["costo_uso"] * p["K_inizio_2012"]).groupby(p["industria_io"]).transform("sum")
    somma_K = p.groupby("industria_io")["K_inizio_2012"].transform("sum")
    p["w"] = p["costo_uso"] * somma_K / somma_cK
    return p[["industria_io", "tipo", "gruppo", "K_inizio_2012", "delta", "r_usato", "costo_uso", "w"]], g.reset_index()


def capitale_capacita(pannello: pd.DataFrame, pesi: pd.DataFrame) -> pd.DataFrame:
    """K^cap_{j,t} (stock di inizio anno) per industria e anno."""
    m = pannello.merge(pesi[["industria_io", "tipo", "w"]], on=["industria_io", "tipo"])
    m = m.dropna(subset=["K_inizio_2012"])
    m["wK"] = m["w"] * m["K_inizio_2012"]
    return m.groupby(["industria_io", "anno"], as_index=False).agg(K_cap=("wK", "sum"), K_somma=("K_inizio_2012", "sum"))


def leggi_g17(percorso, serie: list[str], anni) -> pd.DataFrame:
    """Medie annue delle serie mensili G.17 (formato testo Fed: codice, anno, 12 valori)."""
    righe = []
    with open(percorso, encoding="utf-8", errors="replace") as f:
        for linea in f:
            m = re.match(r'^"([^"]+)"\s+(\d{4})\s+(.*)$', linea.strip())
            if m and m.group(1) in serie and int(m.group(2)) in anni:
                valori = [float(v) for v in m.group(3).split()]
                righe.append({"serie": m.group(1), "anno": int(m.group(2)), "valore": float(np.mean(valori)),
                              "mesi": len(valori)})
    df = pd.DataFrame(righe)
    if df.empty or (df["mesi"] != 12).any():
        raise ValueError("G.17: serie mancanti o anni incompleti")
    return df.pivot(index="serie", columns="anno", values="valore")


def utilizzo_per_industria(u_g17: pd.DataFrame) -> pd.DataFrame:
    righe = []
    for s, inds in G17_IO.items():
        for j in inds:
            for a in u_g17.columns:
                righe.append({"industria_io": j, "anno": a, "serie_g17": s, "u": u_g17.loc[s, a] / 100,
                              "metodo_g17": METODO_G17.get(s, "indagine+capitale")})
    return pd.DataFrame(righe)


def kappa(kcap: pd.DataFrame, x_reale: pd.DataFrame, u: pd.DataFrame, anno: int = 2012) -> pd.DataFrame:
    k = kcap[kcap["anno"] == anno].set_index("industria_io")["K_cap"]
    x = x_reale[x_reale["anno"] == anno].set_index("industria_io")["x_reale"]
    uu = u[u["anno"] == anno].set_index("industria_io")["u"].reindex(k.index).fillna(1.0)
    out = pd.DataFrame({"K_cap_2012": k, "x_2012": x.reindex(k.index), "u_2012": uu})
    out["u_da_g17"] = out.index.isin(u["industria_io"])
    out["kappa"] = out["K_cap_2012"] * out["u_2012"] / out["x_2012"]
    return out.reset_index()


def controllo_g17(kcap: pd.DataFrame, x_reale: pd.DataFrame, kap: pd.DataFrame, u: pd.DataFrame,
                  cap_g17: pd.DataFrame, base: int = 2012) -> pd.DataFrame:
    """Per le industrie G.17: utilizzo implicito u* = x κ / K^cap contro u G.17; crescita della capacità."""
    m = kcap.merge(x_reale, on=["industria_io", "anno"]).merge(kap[["industria_io", "kappa"]], on="industria_io")
    m["u_implicito"] = m["x_reale"] * m["kappa"] / m["K_cap"]
    m = m.merge(u, on=["industria_io", "anno"])
    k_base = m[m["anno"] == base].set_index("industria_io")["K_cap"]
    m["crescita_Kcap"] = m["K_cap"] / m["industria_io"].map(k_base)
    cap = cap_g17.div(cap_g17[base], axis=0)
    m["crescita_cap_g17"] = [cap.loc[s, a] for s, a in zip(m["serie_g17"], m["anno"])]
    m["scarto_u"] = m["u_implicito"] - m["u"]
    return m


def deriva_capacita(ctrl: pd.DataFrame, fine_stima: int = 2014, base: int = 2012) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Deriva annua tra capacità G.17 e capitale per la capacità, per industria.

    θ_j = (crescita capacità G.17 / crescita K^cap)^(1/n) − 1. Con κ_{j,t} = κ_j,2012 / (1+θ_j)^(t−2012)
    la capacità implicita K^cap/κ segue quella G.17 (opzione di calibrazione (a)).
    Verifica fuori campione: θ stimata sul 2012–fine_stima, errore sull'utilizzo negli anni successivi.
    """
    d = ctrl.pivot_table(index="industria_io", columns="anno",
                         values=["crescita_Kcap", "crescita_cap_g17", "u", "u_implicito"])
    metodo = ctrl.drop_duplicates("industria_io").set_index("industria_io")["metodo_g17"]
    ultimo = max(d["u"].columns)
    theta_tutto = (d["crescita_cap_g17"][ultimo] / d["crescita_Kcap"][ultimo]) ** (1 / (ultimo - base)) - 1
    theta_stima = (d["crescita_cap_g17"][fine_stima] / d["crescita_Kcap"][fine_stima]) ** (1 / (fine_stima - base)) - 1
    der = pd.DataFrame({"metodo_g17": metodo, f"deriva_annua_{base}_{ultimo % 100}": theta_tutto,
                        f"deriva_annua_{base}_{fine_stima % 100}": theta_stima}).reset_index()
    righe = []
    for a in [c for c in d["u"].columns if c > fine_stima]:
        senza = (d["u_implicito"][a] - d["u"][a]).abs()
        con = (d["u_implicito"][a] / (1 + theta_stima) ** (a - base) - d["u"][a]).abs()
        for m, idx in metodo.groupby(metodo).groups.items():
            righe.append({"anno": a, "metodo_g17": m, "errore_u_senza_deriva": float(senza[idx].mean()),
                          "errore_u_con_deriva": float(con[idx].mean())})
        righe.append({"anno": a, "metodo_g17": "tutti", "errore_u_senza_deriva": float(senza.mean()),
                      "errore_u_con_deriva": float(con.mean())})
    return der, pd.DataFrame(righe)


def inviluppo_capacita(pan: pd.DataFrame, x: dict, w: pd.Series, industrie: list[str], anni, a0: int) -> pd.DataFrame:
    """Capacità delle industrie senza dati G.17 dall'inviluppo dei massimi storici del rapporto r = x / K^cap (H9c).

    K^cap_{j,t} = Σ_a w_{j,a} K_{j,a,t} (stock di inizio anno a prezzi 2012; per HS: stock residenziale).
    Due metodi, entrambi con la traiettoria osservata ammissibile per costruzione:
      - "inviluppo": capacità = K^cap · max_t r_{j,t} (massimo sul periodo, senza tendenza);
      - "inviluppo_tendenza" (peak-to-peak): log r_{j,t} = α + g·t + e_t; capacità = K^cap · exp(α + g·t + max e),
        cioè la tendenza del rapporto spostata sul massimo storico; deriva θ_j = e^g − 1.
    Restituisce per industria l'utilizzo implicito nell'anno di calibrazione a0 (u = r_a0 / capacità_a0) e θ.
    """
    righe = []
    for j in industrie:
        sub = pan[pan["industria_io"] == j]
        for t in anni:
            s = sub[sub["anno"] == t]
            if j == "HS":
                k = float(s[s["tipo"] == "R"]["K_inizio_2012"].sum())
            else:
                k = float(sum(float(w.get((j, a), 0.0)) * s[s["tipo"] == a]["K_inizio_2012"].sum() for a in ("E", "S", "N")))
            if k > 0 and t in x:
                righe.append({"industria_io": j, "anno": t, "r": float(x[t][j]) / k})
    d = pd.DataFrame(righe)
    out = []
    for j, g in d.groupby("industria_io"):
        g = g.sort_values("anno")
        tt = (g["anno"] - a0).to_numpy(float)
        y = np.log(g["r"].to_numpy(float))
        b = np.polyfit(tt, y, 1)
        e = y - np.polyval(b, tt)
        r0 = float(g.loc[g["anno"] == a0, "r"].iloc[0])
        m_piatto = float(g["r"].max())
        m_tend = float(np.exp(b[1] + e.max()))
        out.append({"industria_io": j, "r_a0": r0, "anno_massimo": int(g.loc[g["r"].idxmax(), "anno"]),
                    "u_inviluppo": r0 / m_piatto, "u_inviluppo_tendenza": r0 / m_tend,
                    "theta_inviluppo_tendenza": float(np.exp(b[0]) - 1), "anno_scarto_massimo": int(g["anno"].iloc[int(e.argmax())])})
    return pd.DataFrame(out)
